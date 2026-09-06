"""Exporter stage: serialise a :class:`ParseResult` to Excel, CSV or JSON.

The Excel layout keeps header fields on top (key/value pairs) followed by a
styled detail table, mirroring how a human would read the original document.
CSV exports the flat detail table (one row per line item) while JSON carries
the full structured payload.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Sequence

from pdf_ocr_toolkit.core.logging import get_logger
from pdf_ocr_toolkit.pipeline.base_parser import ParseResult

logger = get_logger(__name__)

#: Columns used for the flat line-item table (CSV + Excel table section).
ITEM_COLUMNS = ["description", "quantity", "unit_price", "amount"]
ITEM_HEADERS = ["Description", "Quantity", "Unit Price", "Amount"]


def _item_row(item: Dict[str, Any]) -> List[Any]:
    return [item.get("description", ""), item.get("quantity"), item.get("unit_price"), item.get("amount")]


def export_json(result: ParseResult, path: str | Path) -> Path:
    """Write the full structured result as pretty-printed JSON."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(result.to_dict(), indent=2, ensure_ascii=False), encoding="utf-8")
    logger.info("Exported JSON to %s", path)
    return path


def export_csv(result: ParseResult, path: str | Path) -> Path:
    """Write the line-item table as CSV (one row per detail line)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(ITEM_HEADERS)
        for item in result.line_items:
            writer.writerow(_item_row(item.to_dict()))
    logger.info("Exported CSV to %s (%s row(s))", path, len(result.line_items))
    return path


def export_excel(result: ParseResult, path: str | Path) -> Path:
    """Write a human-friendly workbook: header fields + styled detail table."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - openpyxl is a core dep
        raise RuntimeError(
            "openpyxl is required for Excel export. Install it with `pip install openpyxl`."
        ) from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = result.document_type[:31] or "Document"

    title_font = Font(bold=True, size=14, color="1F4E78")
    label_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    header_font = Font(bold=True)
    thin_border = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    # --- Title ---
    sheet.cell(row=1, column=1, value=result.fields.get("_title", result.document_type.title())).font = title_font

    # --- Header fields (key/value) ---
    row = 3
    for key, value in result.fields.items():
        if key.startswith("_"):
            continue
        label_cell = sheet.cell(row=row, column=1, value=key.replace("_", " ").title())
        label_cell.font = label_font
        sheet.cell(row=row, column=2, value="" if value is None else str(value))
        row += 1

    # --- Detail table ---
    row += 1
    header_row = row
    for col, header in enumerate(ITEM_HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    row += 1
    for item in result.line_items:
        for col, value in enumerate(_item_row(item.to_dict()), start=1):
            cell = sheet.cell(row=row, column=col, value=value)
            cell.border = thin_border
            if col in (3, 4) and isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"
            if col == 2 and isinstance(value, (int, float)):
                cell.number_format = "#,##0"
        row += 1

    # --- Column widths ---
    widths = [42, 12, 14, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(path)
    logger.info("Exported Excel to %s (%s line item(s))", path, len(result.line_items))
    return path


def export(result: ParseResult, path: str | Path) -> Path:
    """Dispatch to the right exporter based on the file suffix."""
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return export_excel(result, path)
    if suffix == ".csv":
        return export_csv(result, path)
    if suffix == ".json":
        return export_json(result, path)
    raise ValueError(f"Unsupported export format '{suffix}'. Use .xlsx, .csv or .json.")


def export_many(result: ParseResult, base_path: str | Path, formats: Sequence[str] = ("xlsx",)) -> List[Path]:
    """Export the same result into several formats sharing one base path."""
    base = Path(base_path)
    outputs: List[Path] = []
    for fmt in formats:
        fmt = fmt.lstrip(".")
        outputs.append(export(result, base.with_suffix(f".{fmt}")))
    return outputs
